from fastapi import FastAPI, Request, Body, File, UploadFile, Form, BackgroundTasks
from fastapi.templating import Jinja2Templates
from fastapi.staticfiles import StaticFiles
import pandas as pd
from fastapi.middleware.cors import CORSMiddleware
from fastapi.encoders import jsonable_encoder
import re
from datetime import datetime, date
from fastapi.responses import HTMLResponse
from typing import Optional
from dotenv import load_dotenv
from pydantic import BaseModel
import csv
import os
import io
import docx
import openpyxl
import traceback
import httpx
from langchain_openai import ChatOpenAI

load_dotenv()

app = FastAPI()

# Mount static files and templates
app.mount("/static", StaticFiles(directory="static"), name="static")
templates = Jinja2Templates(directory="templates")

origins = ["*"]

app.add_middleware(
    CORSMiddleware,
    allow_origins=origins,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

@app.get("/", response_class=HTMLResponse)
async def read_root(request: Request):
    """
    Serve the front-end form using Jinja2 templates.
    """
    return templates.TemplateResponse("index.html", {"request": request})

@app.post("/analyze", response_class=HTMLResponse)
async def analyze_cv(
    request: Request,
    name: str = Form(...),
    email: str = Form(...),
    wa_number: str = Form(...),
    education: str = Form(...),
    job_role: str = Form(...),
    cv_file: UploadFile = File(...)
):
    """
    Receive candidate details and a .docx CV, parse the text,
    filter job description, and send it to GLM-4.6 for AI analysis and scoring.
    Finally, append the results to Recruitments.xlsx.
    """
    try:
        # Validate that the file is a docx
        if not cv_file.filename.lower().endswith(".docx"):
            return templates.TemplateResponse("index.html", {
                "request": request,
                "error": "Invalid file format. Please upload a .docx file."
            })

        # Read the file and parse text using python-docx
        file_content = await cv_file.read()
        try:
            doc_stream = io.BytesIO(file_content)
            document = docx.Document(doc_stream)
            cv_text = "\n".join([paragraph.text for paragraph in document.paragraphs])
        except Exception as e:
            return templates.TemplateResponse("index.html", {
                "request": request,
                "error": f"Failed to parse the .docx file. Ensure it is not corrupted. Error: {str(e)}"
            })

        # Load the API key
        api_key = os.getenv("OPENROUTER_API_KEY")
        if not api_key:
            return templates.TemplateResponse("index.html", {
                "request": request,
                "error": "Missing OPENROUTER_API_KEY. Please ensure it is set in your environment variables or .env file."
            })

        # Read job description from Recruitments.xlsx
        excel_path = "Recruitments.xlsx"
        job_description = "Job description not found."
        if os.path.exists(excel_path):
            try:
                df_jobs = pd.read_excel(excel_path, sheet_name='Job Offer')
                # Filter for the selected job role
                job_match = df_jobs[df_jobs['Title Offer'].str.contains(job_role, case=False, na=False)]
                if not job_match.empty:
                    job_description = str(job_match.iloc[0]['Job Description'])
            except Exception as e:
                print(f"Warning: Could not read job description. Error: {e}")
        else:
            print(f"Warning: Excel file {excel_path} not found.")

        # Initialize the OpenRouter LLM using Langchain ChatOpenAI
        http_client = httpx.Client(verify=False)
        llm = ChatOpenAI(
            model="z-ai/glm-4.6",
            temperature=0.3,
            base_url="https://openrouter.ai/api/v1",
            api_key=api_key,
            http_client=http_client,
        )

        # Prepare the prompt for AI Analysis matching Job Description
        prompt = f"""You are an expert HR recruiter and AI Interviewer. Please analyze the following candidate based on their submitted details and their parsed CV text, and evaluate their fit for the requested Job Role.

Candidate Details:
- Name: {name}
- Email: {email}
- WhatsApp: {wa_number}
- Education: {education}
- Applied Job Role: {job_role}

Job Description for {job_role}:
{job_description}

CV Text:
{cv_text}

Task:
1. Provide an overall score for the candidate out of 100 based strictly on how well their qualifications, experience, and skills match the Job Description provided.
2. Provide a detailed analysis of their strengths, weaknesses, and overall fit for this specific role.

Format your response EXACTLY as follows:
Score: [Your Score]
Analysis: [Your detailed analysis]

Don't use any '*' symbol on output. Please strictly use Indonesian language.
"""
        # Invoke the model
        response = llm.invoke(prompt)
        ai_output = response.content

        # Parse the AI response for 'Score' and 'Analysis'
        score = "N/A"
        analysis_text = ai_output

        # Extract the score
        score_match = re.search(r"Score:\s*(\d+)", ai_output, re.IGNORECASE)
        if score_match:
            score = score_match.group(1)

        # Extract just the analysis part
        analysis_match = re.split(r"Analysis:\s*", ai_output, maxsplit=1, flags=re.IGNORECASE)
        if len(analysis_match) > 1:
            analysis_text = analysis_match[1].strip()
        else:
            # Fallback in case the exact format is missed: remove the score line
            analysis_text = re.sub(r"Score:\s*\d+\n?", "", ai_output, flags=re.IGNORECASE).strip()

        # Write results to Recruitments.xlsx in sheet 'Resume Score'
        if os.path.exists(excel_path):
            try:
                # Based on typical column order: Email, Name, Phone, Location (using Education), Score, Resume Text, Analysis
                new_row = pd.DataFrame([{
                    'Email': email,
                    'Name': name,
                    'Phone': wa_number,
                    'Location': education,
                    'Score': int(score) if score.isdigit() else score,
                    'Resume Text': cv_text[:30000],  # Excel cell char limit is 32767
                    'Analysis': analysis_text
                }])

                with pd.ExcelWriter(excel_path, mode='a', engine='openpyxl', if_sheet_exists='overlay') as writer:
                    # Find the last row in the Resume Score sheet
                    if 'Resume Score' in writer.sheets:
                        startrow = writer.sheets['Resume Score'].max_row
                    else:
                        startrow = 0

                    new_row.to_excel(writer, sheet_name='Resume Score', startrow=startrow, index=False, header=False)
            except Exception as e:
                print(f"Warning: Could not save results to excel. Error: {e}")

        # Build the result payload for the template
        result = {
            "name": name,
            "job_role": job_role,
            "score": score,
            "analysis": analysis_text
        }

        return templates.TemplateResponse("index.html", {"request": request, "result": result})

    except Exception as e:
        traceback.print_exc()
        return templates.TemplateResponse("index.html", {
            "request": request,
            "error": f"An unexpected error occurred during processing: {str(e)}\n\nTraceback:\n{traceback.format_exc()}"
        })
